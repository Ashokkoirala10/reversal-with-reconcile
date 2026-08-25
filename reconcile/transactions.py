"""Reads the daily TransactionReport / ibft-transaction .xlsx export — the
same file format the existing reversal_project app consumes."""

from __future__ import annotations

from pathlib import Path

import openpyxl

REQUIRED_COLUMNS = [
    "S NO",
    "Transaction Date",
    "Member Name",
    "Aggregator",
    "Member Transaction Id",
    "Network Reference Id",
    "Session Id",
    "Payment Processor",
    "Transaction Amount",
    "Charge Amount",
    "Debtor Bank",
    "Creditor Bank",
    "Source Message",
    "Overall Status",
]

SOURCE_SHEET_NAME_CANDIDATES = ["Transactions", "transactions"]


class TransactionFileError(Exception):
    pass


def _find_header_row(ws) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True)):
        names = [str(c).strip() if c is not None else "" for c in row]
        if "S NO" in names and "Overall Status" in names:
            col_map = {name: i for i, name in enumerate(names) if name}
            missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
            if missing:
                raise TransactionFileError("The uploaded transaction file is missing column(s): " + ", ".join(missing))
            return row_idx, col_map
    raise TransactionFileError(
        "Could not find the header row (expected a row containing 'S NO' and 'Overall Status'). "
        "Please upload the original TransactionReport / ibft-transaction export."
    )


def load_transactions(path: str | Path) -> list[dict]:
    """Returns a list of dicts, one per transaction row, keyed by the
    original column names."""
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = None
        for name in SOURCE_SHEET_NAME_CANDIDATES:
            if name in wb.sheetnames:
                ws = wb[name]
                break
        if ws is None:
            ws = wb[wb.sheetnames[0]]

        header_row_idx, col_map = _find_header_row(ws)
        rows = []
        for raw in ws.iter_rows(min_row=header_row_idx + 2, values_only=True):
            if raw is None or all(v is None for v in raw):
                continue
            if raw[col_map["S NO"]] is None and all(v is None for v in raw):
                continue
            rows.append({name: (raw[idx] if idx < len(raw) else None) for name, idx in col_map.items()})
        return rows
    finally:
        wb.close()
