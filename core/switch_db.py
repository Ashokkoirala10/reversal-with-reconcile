"""Read-only connector to the switch's own transactional database
(Postgres) — lets a run start from "Fetch from DB" instead of manually
downloading and uploading the exported TransactionReport / ibft-transaction
.xlsx (see core/forms.py:DbFetchForm, core/views.py:upload_view,
reconcile/views.py:reconcile_view).

Deliberately NOT registered as a Django DATABASES alias: this is a schema
we don't own and must never migrate or write to, so it's queried with a
plain psycopg connection instead — that keeps `manage.py migrate`/`test`
from ever touching it.

Connection settings come from SWITCH_DB_* in .env (see settings.SWITCH_DB).
"""

from __future__ import annotations

import datetime as _dt
import logging
from io import BytesIO

import openpyxl
import psycopg
from django.conf import settings
from django.utils import timezone
from openpyxl.utils import get_column_letter

from .services import REQUIRED_COLUMNS

logger = logging.getLogger(__name__)


class SwitchDBError(Exception):
    pass


_BASE36_ALPHABET = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
# The switch schema has no stored "Network Reference Id" column anywhere —
# it's derived from the transaction's own primary key (transaction_entry.id),
# base36-encoded and zero-padded to 12 characters.
_NETWORK_REF_ID_LENGTH = 12

# Mirrors the query confirmed against the UAT switch DB (smartpaycore),
# filtered to one [start, end) window. Kept as an INNER JOIN on
# member_configuration/bank_details like the original — every
# transaction_entry row has a matching member + bank row (verified: the
# joined count equals the unjoined count), so this doesn't silently drop
# rows the way a bad join could.
#
# The "latest payment_status per transaction" lookup below is a LATERAL
# join, not a `SELECT DISTINCT ON (transaction_entry_id) ... GROUP/ORDER`
# subquery — that alternative isn't correlated to the outer date filter,
# so Postgres has to sort/deduplicate the *entire* transaction_payment_status
# table before joining it down to the (usually tiny) date-filtered result.
# That's invisible on a small dataset (fine on UAT, ~12k rows) but turns
# into a multi-minute query once that table has millions of rows (measured
# on the production DB: DISTINCT ON plan ~349k cost / times out vs. this
# LATERAL form's ~36k cost / ~1s for a full day, since it lets Postgres do
# an indexed per-row lookup via idx_tps_transaction_entry_id_desc instead.
_QUERY = """
    SELECT
        t.id AS transaction_id,
        t.created AS transaction_date,
        t.client_unique_id AS member_transaction_id,
        t.transaction_amount AS amount,
        t.fee_amount AS charge_amount,
        t.session_id AS session_id,
        t.payment_processor AS payment_processor,
        bc.bank_name AS creditor_bank_name,
        b.bank_name AS debtor_bank_name,
        m.member_name AS member_name,
        m.aggregator AS aggregator,
        dbs.code AS debit_status,
        dbs.response_unique_id AS debit_response_code,
        cbs.code AS credit_status,
        cbs.response_unique_id AS credit_response_code,
        dbs.message AS source_message,
        cbs.message AS destination_message,
        t.credit_bank_account_number AS credit_account_number,
        t.debit_account_name AS debtor_account_name,
        t.credit_account_name AS creditor_account_name,
        t.debit_bank_account_number AS debit_account_number,
        tps.payment_status AS overall_status
    FROM transaction_entry t
    JOIN member_configuration m ON m.id = t.member_configuration_id
    JOIN bank_details bc ON bc.id = t.creditor_bank_id
    JOIN bank_details b ON b.id = t.debtor_bank_id
    LEFT JOIN LATERAL (
        SELECT tps2.payment_status
        FROM transaction_payment_status tps2
        WHERE tps2.transaction_entry_id = t.id
        ORDER BY tps2.id DESC
        LIMIT 1
    ) tps ON true
    LEFT JOIN transaction_status dbs
        ON dbs.transaction_entry_id = t.id AND dbs.entry_type = 'DR'
    LEFT JOIN transaction_status cbs
        ON cbs.transaction_entry_id = t.id AND cbs.entry_type = 'CR'
    WHERE t.created >= %(start)s AND t.created < %(end)s
    ORDER BY t.id
"""

# Same "must be written as text, never a number" set core/services.py
# enforces on every sheet it generates (see _ID_LIKE_COLUMN_NAMES there) —
# applied here too so a DB-fetched workbook round-trips through
# core.services._load_transactions() / reconcile.transactions.load_transactions()
# exactly like a hand-uploaded one (otherwise a long numeric id can silently
# lose trailing digits once Excel stores it as a number).
_ID_LIKE_COLUMN_NAMES = {
    "Member Transaction Id",
    "Network Reference Id",
    "Session Id",
    "Debit Account Number",
    "Credit Account Number",
}


def _network_reference_id(transaction_id: int) -> str:
    n = int(transaction_id)
    if n == 0:
        digits = "0"
    else:
        chars = []
        while n:
            n, r = divmod(n, 36)
            chars.append(_BASE36_ALPHABET[r])
        digits = "".join(reversed(chars))
    return digits.zfill(_NETWORK_REF_ID_LENGTH)


def _connect():
    cfg = settings.SWITCH_DB
    if not cfg.get("HOST") or not cfg.get("NAME"):
        raise SwitchDBError("Switch DB connection is not configured — set SWITCH_DB_* in .env.")
    try:
        conn = psycopg.connect(
            host=cfg["HOST"],
            port=cfg.get("PORT") or 5432,
            dbname=cfg["NAME"],
            user=cfg.get("USER") or "",
            password=cfg.get("PASSWORD") or "",
            connect_timeout=10,
        )
    except psycopg.OperationalError as exc:
        logger.error("Could not connect to the switch database at %s:%s/%s", cfg.get("HOST"), cfg.get("PORT"), cfg.get("NAME"), exc_info=True)
        raise SwitchDBError(f"Could not connect to the switch database: {exc}") from exc
    # Belt-and-braces: this module only ever issues the one SELECT below,
    # but flipping the session itself to read-only means Postgres rejects
    # any write outright (rather than relying on that staying true by
    # convention) if a future change ever tried to add one here.
    conn.read_only = True
    return conn


def fetch_transactions(start: _dt.datetime, end: _dt.datetime) -> list[dict]:
    """Returns a list of dicts, one per transaction in [start, end), keyed
    by the same column names as an uploaded TransactionReport /
    ibft-transaction export (core.services.REQUIRED_COLUMNS). start/end are
    naive datetimes compared directly against the switch DB's naive
    `created` column (see fetch_ibft_export() for how those are built)."""
    try:
        with _connect() as conn, conn.cursor() as cur:
            cur.execute(_QUERY, {"start": start, "end": end})
            colnames = [d.name for d in cur.description]
            raw_rows = [dict(zip(colnames, row)) for row in cur.fetchall()]
    except psycopg.Error as exc:
        logger.error("Switch database query failed for window %s - %s", start, end, exc_info=True)
        raise SwitchDBError(f"Switch database query failed: {exc}") from exc

    logger.info("Fetched %d transaction(s) from switch DB for window %s - %s", len(raw_rows), start, end)
    rows = []
    for i, r in enumerate(raw_rows, start=1):
        amount = r["amount"]
        charge_amount = r["charge_amount"]
        rows.append(
            {
                "S NO": i,
                "Transaction Date": r["transaction_date"],
                "Member Name": r["member_name"],
                "Aggregator": r["aggregator"],
                "Member Transaction Id": r["member_transaction_id"],
                "Network Reference Id": _network_reference_id(r["transaction_id"]),
                "Session Id": r["session_id"],
                "Payment Processor": r["payment_processor"],
                "Transaction Amount": float(amount) if amount is not None else None,
                "Charge Amount": float(charge_amount) if charge_amount is not None else None,
                "Debtor Bank": r["debtor_bank_name"],
                "Debit Status": r["debit_status"],
                "Debit Response Code": r["debit_response_code"],
                "Debit Account Number": r["debit_account_number"],
                "Debitor Account Name": r["debtor_account_name"],
                "Creditor Bank": r["creditor_bank_name"],
                "Credit Status": r["credit_status"],
                "Credit Response Code": r["credit_response_code"],
                "Credit Account Number": r["credit_account_number"],
                "Creditor Account Name": r["creditor_account_name"],
                "Source Message": r["source_message"],
                "Destination Message": r["destination_message"],
                "Overall Status": r["overall_status"],
            }
        )
    return rows


def build_transactions_workbook(rows: list[dict]) -> bytes:
    """Builds an .xlsx (single "Transactions" sheet, REQUIRED_COLUMNS
    header row) from DB-fetched rows — the same shape
    core.services._load_transactions() / reconcile.transactions.load_transactions()
    expect from a hand-uploaded export, so every downstream step (reversal
    generation, reconciliation) runs unchanged.

    Forces id-like columns to text the cheap way: convert the value to
    `str` before it's ever written (so openpyxl stores it as a text cell
    to begin with — Excel then never has a reason to reinterpret it as a
    number) and set the *column's* number format once, rather than
    visiting every cell after the fact. Doing this per cell via
    `ws.cell(...)` + `cell.number_format = "@"` (one lookup+restyle per
    id-like column per row) measured at over 4 minutes for ~29k rows
    against production — this column-level version does the same job in
    about a second."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(REQUIRED_COLUMNS)

    id_like_indexes = {i for i, name in enumerate(REQUIRED_COLUMNS) if name in _ID_LIKE_COLUMN_NAMES}
    for row in rows:
        values = []
        for i, name in enumerate(REQUIRED_COLUMNS):
            value = row.get(name)
            if i in id_like_indexes and value is not None:
                value = str(value)
            values.append(value)
        ws.append(values)

    for idx, name in enumerate(REQUIRED_COLUMNS, start=1):
        if name in _ID_LIKE_COLUMN_NAMES:
            ws.column_dimensions[get_column_letter(idx)].number_format = "@"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def date_window(
    from_date: _dt.date, to_date: _dt.date, as_of: _dt.datetime | None = None
) -> tuple[_dt.datetime, _dt.datetime]:
    """[start, end) naive-datetime window for `fetch_transactions`, from
    `from_date` 00:00 through `to_date`.

    `to_date`'s end is time-bounded only when it's today — "up to right
    now", so a same-day query only sees what's actually posted so far.
    Any `to_date` before today is a day that's already fully closed out,
    so it's pulled through its own midnight-to-midnight in full — this is
    what makes a T-1 range (which never touches today) naturally need no
    "up to now" cutoff at all.

    NOTE: transaction_entry.created is stored as a naive timestamp in the
    switch DB, holding Asia/Kathmandu wall-clock time (this app's own
    TIME_ZONE) rather than UTC — confirmed against production, not a
    guess."""
    start = _dt.datetime.combine(from_date, _dt.time.min)
    if as_of is not None:
        end = as_of
    elif to_date >= timezone.localdate():
        end = timezone.localtime(timezone.now()).replace(tzinfo=None)
    else:
        end = _dt.datetime.combine(to_date + _dt.timedelta(days=1), _dt.time.min)
    return start, end


def fetch_ibft_export(
    from_date: _dt.date, to_date: _dt.date, as_of: _dt.datetime | None = None
) -> tuple[bytes, int]:
    """High-level entry point behind "Fetch from DB": pulls every
    transaction from `from_date` 00:00 through `to_date`, and returns
    (workbook_bytes, row_count). Raises SwitchDBError if nothing was
    found, so callers can show that the same way an empty/bad upload
    would be reported."""
    start, end = date_window(from_date, to_date, as_of)
    rows = fetch_transactions(start, end)
    if not rows:
        raise SwitchDBError(
            f"No transactions found in the switch database for {from_date} through {to_date} (up to {end:%Y-%m-%d %H:%M})."
        )
    return build_transactions_workbook(rows), len(rows)
